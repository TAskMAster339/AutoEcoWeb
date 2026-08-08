"""Импорт/экспорт данных: парсинг универсального формата, round-trip, сервис."""

import datetime
import io
from collections.abc import Sequence
from decimal import Decimal

import openpyxl
import pytest
from src.models.tag import Tag
from src.models.transaction import Transaction
from src.models.user import User
from src.repositories.tag import TagRepository
from src.repositories.transaction import TransactionRepository
from src.repositories.user import UserRepository
from src.schemas.import_export import ImportRowIn
from src.services.import_export import (
    ExportRow,
    ImportExportService,
    ImportFormatError,
    build_export_workbook,
    parse_import_file,
)

HEADER = ["Дата", "Категория", "Магазин", "Описание", "Доход", "Расход"]


def _make_xlsx(rows: list[Sequence[object]], sheet: str = "Данные") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_csv(rows: list[list[object]], encoding: str = "utf-8") -> bytes:
    lines = [";".join("" if cell is None else str(cell) for cell in row) for row in rows]
    return "\n".join(lines).encode(encoding)


async def _make_user(session) -> User:
    return await UserRepository(session).create(
        email="import@test.ru",
        password_hash="x" * 60,
    )


def _service(session) -> ImportExportService:
    return ImportExportService(
        session,
        TransactionRepository(session),
        TagRepository(session),
    )


# ---------- парсинг xlsx ----------


def test_parse_xlsx_basic():
    content = _make_xlsx(
        [
            HEADER,
            [datetime.datetime(2025, 9, 21, 12, 0), "Еда", "Пятёрочка", "Молоко", None, 89.9],
            ["22.09.2025", "ЗП", "", "Аванс", "40000", None],
            ["23.09.2025", None, None, "Шаурма", None, "250,50"],
        ]
    )
    preview = parse_import_file("test.xlsx", content)
    assert preview.total == 3
    assert preview.valid == 3
    assert preview.invalid == 0
    rows = preview.rows
    assert rows[0].date == datetime.date(2025, 9, 21)
    assert rows[0].category == "Еда"
    assert rows[0].store == "Пятёрочка"
    assert rows[0].income == Decimal("0")
    assert rows[0].expense == Decimal("89.90")
    assert rows[1].date == datetime.date(2025, 9, 22)
    assert rows[1].income == Decimal("40000.00")
    assert rows[2].store is None
    assert rows[2].expense == Decimal("250.50")


def test_parse_xlsx_header_not_first_row_and_junk_columns():
    content = _make_xlsx(
        [
            ["Отчёт за сентябрь"],
            ["", "", "", "", "", ""],
            HEADER + ["Мусор", "Итого"],
            ["24.09.2025", "Транспорт", "", "Метро", None, 60, 60, 60],
        ]
    )
    preview = parse_import_file("x.xlsx", content)
    assert preview.total == 1
    assert preview.rows[0].description == "Метро"
    assert preview.rows[0].expense == Decimal("60.00")


def test_parse_xlsx_row_errors():
    content = _make_xlsx(
        [
            HEADER,
            ["недата", "Еда", "", "", None, 10],
            ["25.09.2025", "Еда", "", "", None, 10],  # пустое описание
            ["26.09.2025", "Еда", "", "И то и то", 5, 10],
            ["27.09.2025", "Еда", "", "Нет суммы", None, None],
            ["28.09.2025", "Еда", "", "Ок", None, "10,5"],
        ]
    )
    preview = parse_import_file("x.xlsx", content)
    assert preview.total == 5
    assert preview.valid == 1
    assert preview.invalid == 4
    assert "Некорректная дата" in preview.rows[0].errors
    assert "Пустое описание" in preview.rows[1].errors
    assert "Заполнены и Доход, и Расход" in preview.rows[2].errors
    assert "Нет суммы" in preview.rows[3].errors[0]
    assert preview.rows[4].errors == []


def test_parse_xlsx_blank_rows_skipped():
    content = _make_xlsx([HEADER, [None, None, None, None, None, None], ["01.10.2025", "", "", "Кофе", None, 150]])
    preview = parse_import_file("x.xlsx", content)
    assert preview.total == 1


def test_parse_csv_semicolon_and_cp1251():
    rows = [
        HEADER,
        ["01.10.2025", "Еда", "Магнит", "Хлеб", None, "45,50"],
        ["02.10.2025", "ЗП", "", "Зарплата", "80 000", None],
    ]
    content = _make_csv(rows, encoding="cp1251")
    preview = parse_import_file("отчёт.csv", content)
    assert preview.total == 2
    assert preview.valid == 2
    assert preview.rows[0].expense == Decimal("45.50")
    assert preview.rows[1].income == Decimal("80000.00")


def test_parse_invalid_file_raises():
    with pytest.raises(ImportFormatError):
        parse_import_file("data.txt", b"hello world")


def test_parse_no_header_raises():
    content = _make_xlsx([["a", "b"], ["1", "2"]])
    with pytest.raises(ImportFormatError):
        parse_import_file("x.xlsx", content)


# ---------- сборка и round-trip экспорта ----------


def test_build_export_workbook_roundtrip():
    rows = [
        ExportRow(
            date=datetime.date(2025, 9, 21),
            category="Еда",
            store="Пятёрочка",
            description="Молоко",
            income=Decimal("0"),
            expense=Decimal("89.90"),
        ),
        ExportRow(
            date=datetime.date(2025, 10, 1),
            category="ЗП",
            store=None,
            description="Аванс",
            income=Decimal("40000.00"),
            expense=Decimal("0"),
        ),
    ]
    content = build_export_workbook(rows)
    preview = parse_import_file("autoeco-export.xlsx", content)
    assert preview.total == 2
    assert preview.valid == 2
    by_desc = {row.description: row for row in preview.rows}
    assert by_desc["Молоко"].date == datetime.date(2025, 9, 21)
    assert by_desc["Молоко"].store == "Пятёрочка"
    assert by_desc["Молоко"].expense == Decimal("89.90")
    assert by_desc["Аванс"].income == Decimal("40000.00")
    assert by_desc["Аванс"].category == "ЗП"


def test_build_export_workbook_sheets_by_month():
    rows = [
        ExportRow(datetime.date(2025, 9, 5), None, None, "A", Decimal("0"), Decimal("1")),
        ExportRow(datetime.date(2025, 10, 5), None, None, "B", Decimal("0"), Decimal("2")),
    ]
    content = build_export_workbook(rows)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Сентябрь 2025", "Октябрь 2025"]
    assert list(wb["Сентябрь 2025"].iter_rows(values_only=True))[0] == tuple(HEADER)


def test_build_export_workbook_empty():
    content = build_export_workbook([])
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Транзакции"]
    assert list(wb["Транзакции"].iter_rows(values_only=True))[0] == tuple(HEADER)


# ---------- сервис: импорт ----------


def _row(**overrides) -> ImportRowIn:
    base = {
        "date": datetime.date(2025, 9, 21),
        "category": None,
        "store": None,
        "description": "Молоко",
        "income": Decimal("0"),
        "expense": Decimal("89.90"),
    }
    base.update(overrides)
    return ImportRowIn(**base)


async def test_import_rows_creates_tags_and_maps_fields(session):
    user = await _make_user(session)
    result = await _service(session).import_rows(
        user,
        [
            _row(description="Молоко", category="Еда", store="Пятёрочка"),
            _row(
                date=datetime.date(2025, 9, 22),
                description="Аванс",
                category="ЗП",
                income=Decimal("40000.00"),
                expense=Decimal("0"),
            ),
        ],
    )
    assert result.imported == 2
    assert result.tags_created == ["Еда", "ЗП"]

    tags = await TagRepository(session).list_all(user.id)
    assert {t.name for t in tags} == {"Еда", "ЗП"}

    txs = await TransactionRepository(session).list_all(user.id)
    assert len(txs) == 2
    tx, seller = txs[0]
    assert tx.user_id == user.id
    assert tx.receipt_id is None
    assert tx.name == "Молоко"
    assert tx.normalized_name == "молоко"
    assert tx.seller_name == "Пятёрочка"
    assert tx.amount == Decimal("89.90")
    assert tx.operation_type == 1  # расход
    assert tx.check_datetime.date() == datetime.date(2025, 9, 21)
    assert tx.tag_id is not None

    income_tx, _ = txs[1]
    assert income_tx.amount == Decimal("40000.00")
    assert income_tx.operation_type == 2  # доход


async def test_import_rows_reuses_existing_tag(session):
    user = await _make_user(session)
    tag = await TagRepository(session).create(
        user_id=user.id,
        name="Еда",
        color="#3B82F6",
        icon=None,
    )
    result = await _service(session).import_rows(user, [_row(category="Еда")])
    assert result.tags_created == []
    tags = await TagRepository(session).list_all(user.id)
    assert len(tags) == 1
    assert tags[0].id == tag.id


async def test_import_rows_empty(session):
    user = await _make_user(session)
    result = await _service(session).import_rows(user, [])
    assert result.imported == 0
    assert result.tags_created == []


async def test_import_rows_standalone_are_visible_in_export(session):
    user = await _make_user(session)
    await _service(session).import_rows(
        user,
        [
            _row(description="Молоко", category="Еда", store="Пятёрочка"),
            _row(
                date=datetime.date(2025, 9, 22),
                description="Аванс",
                income=Decimal("40000.00"),
                expense=Decimal("0"),
            ),
        ],
    )
    rows = await _service(session).export_rows(user)
    assert len(rows) == 2
    by_desc = {row.description: row for row in rows}
    assert by_desc["Молоко"].category == "Еда"
    assert by_desc["Молоко"].store == "Пятёрочка"
    assert by_desc["Молоко"].expense == Decimal("89.90")
    assert by_desc["Аванс"].income == Decimal("40000.00")
    assert by_desc["Аванс"].store is None


async def test_export_rows_excludes_tag_field_when_no_tag(session):
    user = await _make_user(session)
    await _service(session).import_rows(user, [_row(description="Без тега")])
    rows = await _service(session).export_rows(user)
    assert rows[0].category is None
