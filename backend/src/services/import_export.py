"""Импорт/экспорт данных пользователя (универсальный формат).

Канонический формат (источник истины — экспорт приложения):

    Дата | Категория | Магазин | Описание | Количество | Единица |
    Цена | Комментарий | Доход | Расход

- Файлы: .xlsx (любое число листов, у каждого шапка в первых 5 строках)
  или .csv (UTF-8/CP1251, разделитель ';' или ',').
- Порядок колонок не важен, лишние колонки игнорируются.
- Никакой эвристики «под конкретный файл»: строка = одна транзакция,
  значение = как в файле. Ошибки строки собираются в preview и отдаются
  пользователю до импорта.

Чистые функции (парсинг/сборка) не трогают БД и покрыты тестами.
"""  # noqa: RUF002

from __future__ import annotations

import csv
import datetime
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Literal
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.ext.asyncio import AsyncSession
from src.models.tag import Tag
from src.models.transaction import Transaction
from src.models.user import User
from src.repositories.alias import AliasRepository
from src.repositories.tag import TagRepository
from src.repositories.transaction import TransactionRepository
from src.schemas.import_export import (
    ImportPreview,
    ImportResult,
    ImportRowIn,
    ImportRowPreview,
)
from src.services.aliases import AliasService
from src.services.receipt_parser import normalize_product_name
from src.services.sellers import SellerService

__all__ = [
    "ExportRow",
    "ImportExportService",
    "ImportFormatError",
    "build_export_workbook",
    "parse_import_file",
]

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 МБ


_TAG_PALETTE = [
    "#16A34A",
    "#3B82F6",
    "#8B5CF6",
    "#F97316",
    "#06B6D4",
    "#F59E0B",
    "#EF4444",
    "#65A30D",
]

_RU_MONTHS = [
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
]

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("дата", "date", "дата операции"),
    "category": ("категория", "category", "тег", "tag"),
    "store": ("магазин", "store", "продавец", "seller", "место"),
    "description": (
        "описание",
        "description",
        "название",
        "name",
    ),
    "quantity": ("количество", "quantity", "кол-во", "qty"),
    "unit": ("единица", "unit", "ед. изм.", "единица измерения"),
    "price": ("цена", "price", "цена за единицу"),
    "comment": ("комментарий", "comment", "примечание", "note"),
    "income": ("доход", "income", "приход"),
    "expense": ("расход", "expense"),
}

_DATE_DMY_RE = re.compile(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$")


class ImportFormatError(ValueError):
    """Файл не похож на поддерживаемый формат (не xlsx/csv или нет шапки)."""


@dataclass(frozen=True)
class ExportRow:
    """Строка экспорта (и канонический формат файла)."""

    date: datetime.date
    category: str | None
    store: str | None
    description: str
    quantity: Decimal | None
    unit: str | None
    price: Decimal | None
    comment: str | None
    income: Decimal | None
    expense: Decimal | None


def _clean_str(value: object) -> str | None:
    if value is None:
        return None
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return cleaned or None


def _parse_date(value: object) -> datetime.date | None:  # noqa: PLR0911
    """Excel-дата, DD.MM.YYYY, DD.MM.YY, YYYY-MM-DD, ISO-строка."""
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    match = _DATE_DMY_RE.fullmatch(text)
    if match is not None:
        day, month, year = (int(g) for g in match.groups())
        if year < 100:  # noqa: PLR2004
            year += 2000
        try:
            return datetime.date(year, month, day)
        except ValueError:
            return None
    try:
        return datetime.date.fromisoformat(text)
    except ValueError:
        pass
    try:
        return datetime.datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _parse_money(value: object) -> Decimal | None:  # noqa: PLR0911
    """Число или текст '5 000,50' / '5000.5' → Decimal(0.01). None — не сумма."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    if isinstance(value, (int, float)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    text = str(value).strip()
    if not text:
        return None
    text = (
        text.replace("\u00a0", "")
        .replace(" ", "")
        .replace("\u20bd", "")  # ₽
        .replace("руб", "")  # noqa: RUF001
        .replace("р.", "")  # noqa: RUF001
        .replace(",", ".")
    )
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _parse_quantity(value: object) -> Decimal | None:
    """Количество хранится с точностью до трёх знаков, в отличие от денег."""  # noqa: RUF002
    parsed = _parse_money(value)
    if parsed is None:
        return None
    try:
        return Decimal(str(value).strip().replace(",", ".")).quantize(Decimal("0.001"))
    except InvalidOperation:
        return parsed.quantize(Decimal("0.001"))


# ---------- чтение файла ----------


RawSheet = Sequence[Sequence[object]]


def _read_sheets(filename: str, content: bytes) -> list[RawSheet]:
    """Сырые строки по листам. Различает xlsx (openpyxl) и csv (stdlib)."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text: str | None = None
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = content.decode("utf-8", errors="replace")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        rows: list[Sequence[object]] = [
            [cell.strip() if isinstance(cell, str) else cell for cell in row]
            for row in csv.reader(io.StringIO(text), delimiter=delimiter)
        ]
        return [rows]
    if name.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                data_only=True,
                read_only=True,
            )
        except (InvalidFileException, OSError, ValueError) as exc:
            raise ImportFormatError(
                "Не удалось прочитать файл: нужен .xlsx или .csv",  # noqa: RUF001
            ) from exc
        sheets: list[RawSheet] = [
            [list(row) for row in worksheet.iter_rows(values_only=True)]
            for worksheet in workbook.worksheets
        ]
        return sheets
    raise ImportFormatError("Поддерживаются только файлы .xlsx и .csv")


def _header_map(row: Sequence[object]) -> dict[str, int] | None:
    """Колонки по именам заголовков. Нужны: дата + описание + хотя бы одна сумма."""
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        name = _clean_str(cell)
        if not name:
            continue
        normalized = name.lower()
        for field, aliases in _HEADER_ALIASES.items():
            if normalized in aliases and field not in mapping:
                mapping[field] = index
    if {"date", "description"} <= mapping.keys() and (
        "income" in mapping or "expense" in mapping
    ):
        return mapping
    return None


def _cell(row: Sequence[object], mapping: dict[str, int], field: str) -> object:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _row_to_preview(
    index: int,
    row: Sequence[object],
    mapping: dict[str, int],
) -> ImportRowPreview:
    errors: list[str] = []

    parsed_date = _parse_date(_cell(row, mapping, "date"))
    if parsed_date is None:
        errors.append("Некорректная дата")

    category = _clean_str(_cell(row, mapping, "category"))
    store = _clean_str(_cell(row, mapping, "store"))
    description = _clean_str(_cell(row, mapping, "description")) or ""
    if not description:
        errors.append("Пустое описание")

    quantity = _parse_quantity(_cell(row, mapping, "quantity"))
    price = _parse_money(_cell(row, mapping, "price"))
    unit = _clean_str(_cell(row, mapping, "unit"))
    comment = _clean_str(_cell(row, mapping, "comment"))
    if quantity is not None and quantity <= 0:
        errors.append("Количество должно быть больше нуля")
        quantity = None
    if price is not None and price < 0:
        errors.append("Цена не может быть отрицательной")
        price = None

    income = _parse_money(_cell(row, mapping, "income"))
    expense = _parse_money(_cell(row, mapping, "expense"))
    if income is not None and income < 0:
        errors.append("Доход не может быть отрицательным")
        income = Decimal("0")
    if expense is not None and expense < 0:
        errors.append("Расход не может быть отрицательным")
        expense = Decimal("0")
    if income is not None and expense is not None and income > 0 and expense > 0:
        errors.append("Заполнены и Доход, и Расход")
    elif income is None and expense is None:
        errors.append("Нет суммы: заполните Доход или Расход")

    operation_kind = (
        "income"
        if income is not None and (income > 0 or expense is None)
        else "expense"
    )

    return ImportRowPreview(
        index=index,
        date=parsed_date,
        category=category,
        store=store,
        description=description,
        quantity=quantity,
        unit=unit,
        price=price,
        comment=comment,
        income=income or Decimal("0"),
        expense=expense or Decimal("0"),
        operation_kind=operation_kind,
        errors=errors,
    )


def parse_import_file(filename: str, content: bytes) -> ImportPreview:
    """Файл → предпросмотр строк. Чистая функция, БД не трогает."""
    sheets = _read_sheets(filename, content)
    rows: list[ImportRowPreview] = []
    for sheet in sheets:
        header_index: int | None = None
        mapping: dict[str, int] | None = None
        for index, row in enumerate(sheet[:5]):
            candidate = _header_map(row)
            if candidate is not None:
                header_index = index
                mapping = candidate
                break
        if mapping is None:
            continue  # лист без шапки пропускаем целиком
        for row in sheet[header_index + 1 :]:  # type: ignore[operator]
            if all(_clean_str(cell) is None for cell in row):
                continue  # пустая строка
            rows.append(_row_to_preview(len(rows), row, mapping))
    if not rows:
        raise ImportFormatError(
            "В файле нет данных с шапкой «Дата | Категория | Магазин | Описание | "  # noqa: RUF001
            "Доход | Расход»",
        )
    valid = sum(1 for row in rows if not row.errors)
    return ImportPreview(
        rows=rows,
        total=len(rows),
        valid=valid,
        invalid=len(rows) - valid,
    )


# ---------- сборка экспорта ----------


_EXPORT_HEADER = [
    "Дата",
    "Категория",
    "Магазин",
    "Описание",
    "Количество",
    "Единица",
    "Цена",
    "Комментарий",
    "Доход",
    "Расход",
]
_EXPORT_WIDTHS = (12, 22, 22, 48, 14, 12, 14, 36, 14, 14)


def _fill_export_sheet(worksheet, rows: list[ExportRow]) -> None:
    worksheet.append(_EXPORT_HEADER)
    for item in sorted(rows, key=lambda row: row.date):
        worksheet.append(
            [
                item.date,
                item.category or "",
                item.store or "",
                item.description,
                item.quantity,
                item.unit or "",
                item.price,
                item.comment or "",
                item.income,
                item.expense,
            ],
        )
    for row_cells in worksheet.iter_rows(min_row=2):
        row_cells[0].number_format = "DD.MM.YYYY"
        row_cells[4].number_format = "#,##0.000"
        row_cells[6].number_format = "#,##0.00"
        row_cells[8].number_format = "#,##0.00"
        row_cells[9].number_format = "#,##0.00"
    worksheet.freeze_panes = "A2"
    for index, width in enumerate(_EXPORT_WIDTHS, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width


def build_export_workbook(
    rows: list[ExportRow],
    *,
    layout: Literal["monthly", "single"] = "monthly",
) -> bytes:
    """Строки → .xlsx: один общий лист или отдельный лист на каждый месяц."""
    workbook = Workbook()
    if workbook.active is not None:
        workbook.remove(workbook.active)

    if layout == "single":
        _fill_export_sheet(workbook.create_sheet("Транзакции"), rows)
    else:
        by_month: dict[tuple[int, int], list[ExportRow]] = {}
        for row in rows:
            by_month.setdefault((row.date.year, row.date.month), []).append(row)
        if not by_month:
            _fill_export_sheet(workbook.create_sheet("Транзакции"), [])
        for (year, month), items in sorted(by_month.items()):
            _fill_export_sheet(
                workbook.create_sheet(f"{_RU_MONTHS[month - 1]} {year}"),
                items,
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------- сервис (БД) ----------


def _tag_color(name: str) -> str:
    return _TAG_PALETTE[sum(ord(ch) for ch in name) % len(_TAG_PALETTE)]


class ImportExportService:
    """Импорт подтверждённых строк и сборка данных для экспорта."""

    def __init__(
        self,
        session: AsyncSession,
        tx_repo: TransactionRepository,
        tag_repo: TagRepository,
        alias_repo: AliasRepository,
        seller_service: SellerService,
    ) -> None:
        self._session = session
        self._tx_repo = tx_repo
        self._tag_repo = tag_repo
        self._alias_repo = alias_repo
        self._seller_service = seller_service

    async def import_rows(self, user: User, rows: list[ImportRowIn]) -> ImportResult:
        """Валидированные строки → теги (авто-создание) + bulk-insert транзакций."""
        existing: dict[str, Tag] = {
            tag.name: tag for tag in await self._tag_repo.list_all(user.id)
        }
        tags_created: list[str] = []
        transactions: list[Transaction] = []
        product_aliases = (
            await self._alias_repo.list_all(user.id, scope="product")
            if self._alias_repo is not None
            else []
        )

        for row in rows:
            tag_id: UUID | None = None
            if row.category:
                tag = existing.get(row.category)
                if tag is None:
                    tag = await self._tag_repo.create(
                        user_id=user.id,
                        name=row.category,
                        color=_tag_color(row.category),
                        icon=None,
                    )
                    existing[row.category] = tag
                    tags_created.append(row.category)
                tag_id = tag.id

            is_income = row.operation_kind == "income" or (
                row.operation_kind is None and row.income > 0
            )
            amount = row.income if is_income else row.expense
            product_resolved = AliasService.resolve_with_alias(
                product_aliases,
                row.description,
            )
            seller = (
                await self._seller_service.get_or_create(user.id, row.store)
                if row.store and self._seller_service is not None
                else None
            )
            transactions.append(
                Transaction(
                    user_id=user.id,
                    receipt_id=None,
                    position=None,
                    name=row.description,
                    normalized_name=normalize_product_name(product_resolved.value),
                    name_alias_id=product_resolved.alias.id
                    if product_resolved.alias is not None
                    else None,
                    seller_id=seller.id if seller is not None else None,
                    quantity=row.quantity,
                    unit=row.unit,
                    price=row.price,
                    amount=amount.quantize(Decimal("0.01")),
                    comment=row.comment,
                    operation_type=2 if is_income else 1,
                    check_datetime=datetime.datetime.combine(
                        row.date,
                        datetime.time(12, 0),
                        tzinfo=datetime.timezone.utc,
                    ),
                    tag_id=tag_id,
                ),
            )

        if transactions:
            await self._tx_repo.create_many_standalone(transactions)
        return ImportResult(
            imported=len(transactions),
            errors=[],
            tags_created=tags_created,
        )

    async def export_rows(self, user: User) -> list[ExportRow]:
        """Все транзакции пользователя (включая из чеков) → строки канонического формата."""  # noqa: E501, RUF002
        tags = {tag.id: tag.name for tag in await self._tag_repo.list_all(user.id)}
        rows: list[ExportRow] = []
        for tx, seller_name in await self._tx_repo.list_all(user.id):
            is_income = tx.operation_type in (2, 3)
            rows.append(
                ExportRow(
                    date=tx.check_datetime.date(),
                    category=tags.get(tx.tag_id) if tx.tag_id else None,
                    # list_all already returns the effective (alias-resolved)
                    # seller value, including the receipt fallback.
                    store=seller_name,
                    description=tx.name,
                    quantity=tx.quantity,
                    unit=tx.unit,
                    price=tx.price,
                    comment=tx.comment,
                    income=tx.amount if is_income else None,
                    expense=tx.amount if not is_income else None,
                ),
            )
        return rows
